#!/usr/bin/env python3
"""
Documents the editor approved in Pulse join the Object Matrix (T-0111, the editor's rule, 2026-10-30: "Only
things I have approved"). Runs at publish, after the Backstory draft has shipped.

Reads  editorial/research-ledger.json       finds with approved_by "editor" (Pulse writes this file)
       digest/data/backstory.json           which story approved each one, and that story's category
Writes editorial/NTK_Backstory_Object_Matrix.xlsx   one new row per document (Objects sheet)
       editorial/matrix-additions.json      what this step has handled, so each find is added once. Its own
                                            file, never the ledger: Pulse writes the ledger at any moment,
                                            and a publish that also wrote it could fail to push.

A news report is never added: a matrix link is a primary source (docs/decisions.md, 2026-10-02 and 2026-10-13).
A document whose link is already in the matrix is not added twice. NTK Row is the story's category when that is
one of the fixed rows, and blank otherwise (a document outside the American rows, as 2026-10-06 anticipated).
The pool Pulse reads is rebuilt afterwards by editorial/export_for_pulse.py.

  python3 editorial/graduate_finds.py [--dry-run]
"""
import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
LEDGER = ROOT / "editorial" / "research-ledger.json"
DONE = ROOT / "editorial" / "matrix-additions.json"
MATRIX = ROOT / "editorial" / "NTK_Backstory_Object_Matrix.xlsx"
BACKSTORY = ROOT / "digest" / "data" / "backstory.json"
NEWS = ("nytimes.com", "washingtonpost.com", "wsj.com", "apnews.com", "reuters.com", "bloomberg.com", "politico.com",
        "cnn.com", "nbcnews.com", "cbsnews.com", "foxnews.com", "theguardian.com", "bbc.co", "bbc.com", "npr.org")


def row_for(f, data):
    """The matrix row: the category's title when it is a fixed row, and the story's scope as the theme."""
    rows = {r["id"]: r for r in data.get("rows", [])}
    for p in data.get("todays_pairings", []):
        if p.get("story_id") == f.get("approved_for") or p.get("story_id") == f.get("story_id"):
            r = rows.get(p.get("row")) or {}
            fixed = r.get("title") if r and r.get("stratum") != "provisional" else None
            return fixed, (p.get("composed") or {}).get("scope") or f.get("researched_for") or ""
    return None, f.get("researched_for") or ""


def matrix_row(f, ntk_row, theme):
    year = str(f.get("year") or "")[:4]
    date = str(f.get("date") or "")
    dm = ""
    if len(date) == 10:
        d = datetime.strptime(date, "%Y-%m-%d")
        dm = f"{d.day} {d.strftime('%B')}"
    y = int(year) if year.isdigit() else None
    return [y // 10 * 10 if y else None, year, dm or None, date or year, theme or None, f.get("title"),
            f.get("author") or None, None, None, f.get("source") or None, ntk_row, None,
            ("Addition, pre-1981" if y and y < 1981 else "Extension 1981-2026"),
            "researched" if date else "year only", f.get("url")]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--ledger", default=str(LEDGER))
    ap.add_argument("--matrix", default=str(MATRIX))
    ap.add_argument("--backstory", default=str(BACKSTORY))
    ap.add_argument("--done", default=str(DONE))
    a = ap.parse_args()
    if not Path(a.ledger).exists():
        print("no research ledger; nothing to add")
        return
    ledger = json.loads(Path(a.ledger).read_text())
    done = json.loads(Path(a.done).read_text()) if Path(a.done).exists() else {"handled": {}}
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    data = json.loads(Path(a.backstory).read_text()) if Path(a.backstory).exists() else {}
    wb = openpyxl.load_workbook(a.matrix)
    sh = wb["Objects"]
    url_col = [c.value for c in sh[1]].index("Source URL") + 1
    have = {str(sh.cell(row=i, column=url_col).value or "").strip() for i in range(2, sh.max_row + 1)}
    added, skipped = [], []
    for f in ledger.get("finds", []):
        if f.get("approved_by") != "editor" or f.get("id") in done["handled"]:
            continue
        url = str(f.get("url") or "").strip()
        host = url.split("/")[2].lower() if url.count("/") >= 2 else ""
        if not url.startswith("http") or f.get("source_kind") == "news" or any(host == n or host.endswith("." + n) for n in NEWS):
            skipped.append((f.get("title"), "news report or no link: never enters the matrix"))
            done["handled"][f["id"]] = {"at": now, "result": "not eligible", "title": f.get("title"), "url": url}
            continue
        if url in have:
            skipped.append((f.get("title"), "already in the matrix"))
            done["handled"][f["id"]] = {"at": now, "result": "already in the matrix", "title": f.get("title"), "url": url}
            continue
        ntk_row, theme = row_for(f, data)
        sh.append(matrix_row(f, ntk_row, theme))
        have.add(url)
        done["handled"][f["id"]] = {"at": now, "result": "added", "title": f.get("title"), "url": url, "ntk_row": ntk_row}
        added.append((f.get("year"), f.get("title"), ntk_row or "(no fixed row)"))
    for y, t, r in added:
        print(f"  added  {y} {str(t)[:60]}  [{r}]")
    for t, why in skipped:
        print(f"  skip   {str(t)[:60]}: {why}")
    print(f"{len(added)} document(s) added to the Object Matrix, {len(skipped)} skipped")
    if a.dry_run or not (added or skipped):
        return
    if added:
        wb.save(a.matrix)
    Path(a.done).write_text(json.dumps(done, indent=1, ensure_ascii=False) + "\n")


if __name__ == "__main__":
    sys.exit(main())
